import os
import sys
import time
import logging
import pandas as pd
from pathlib import Path
from docx import Document
from docx.shared import Pt
from dotenv import load_dotenv
from watchdog.observers import Observer
from watchdog.events import FileSystemEventHandler

logging.basicConfig(level=logging.INFO, format="> %(message)s at %(asctime)s.", datefmt="%H:%M:%S")

class Main(FileSystemEventHandler):
    def on_created(self, event): 
        path = Path(event.src_path)
        is_file = path.is_file()
        is_csv = path.suffix == '.csv'
        is_primary = '_2' not in path.stem

        if is_file and is_csv and is_primary and self._is_file_loaded(path):
            self._reporting(path.stem, self._processing(pd.read_csv(path, skiprows=[0,2], index_col=0)))

    def _reporting(self, name, processed_data):
        # print(processed_data, type(processed_data), flush=True)

        if '_1' in name:
            name = f'{name.split("_1")[0]}.docx'
        else:
            name = f'{name}.docx'

        try:
            report = Document(Path(REPORTS_DIRECTORY) / name)

            for row in report.tables[1].rows:
                for cell in row.cells:
                    for key, value in processed_data.items():
                        if cell.text == key:
                            cell.text = str(value)
                    
                    if '+' in cell.text and cell.text:
                        cell.text = ''

            report.save(Path(REPORTS_DIRECTORY) / name)
            print('saved', flush=True)
        except (FileNotFoundError, UnicodeDecodeError):
            pass

    def _processing(self, data):    
        key_map = {
            'Initial area at Area reduction': '+AREA',
            'Area': '+AREA',
            'Force at Yield (Offset 0.2 %)': '+YIELD',
            'Tensile stress at Yield (Offset 0.2 %)': '+S2',
            'Tensile stress at Maximum Force': '+S1',
            'Maximum Force': '+ULT',
            'Elongation after fracture': '+E',
            'Reduction of area at Area reduction': '+RA'
        }

        data = data.rename(columns=key_map)
        placeholders = set(key_map.values())
        valid_columns = [column for column in data.columns if column in placeholders]
        data = data[valid_columns]
        data = data.to_dict()
        processed_data = dict()

        for key, value in data.items():
            for sub_key, sub_value in value.items():
                    if key in ['+S1', '+S2']:
                        if sub_value < 50000:
                            sub_value = round(sub_value/100) * 100
                        elif sub_value < 100000:
                            sub_value = round(sub_value/500) * 500
                        else:
                            sub_value = round(sub_value/1000) * 1000

                        processed_data[f'{key}{sub_key}+'] = f'{sub_value:,}'

                    elif key in ['+E', '+RA']:
                        processed_data[f'{key}{sub_key}+'] = round(sub_value)
                    else:
                        processed_data[f'{key}{sub_key}+'] = f'{sub_value:,}'

        return processed_data

    def _is_file_loaded(self, data_file, max_attemps = 20):
        attempts = 0

        while attempts < max_attemps:
            try:
                with open(data_file, 'r+'):
                    return True
            except (PermissionError, FileNotFoundError, IOError):
                attempts += 1
                time.sleep(0.5)
        return False

if __name__ == "__main__":    
    load_dotenv()
    DATA_DIRECTORY = os.getenv("DATA_DIRECTORY")
    REPORTS_DIRECTORY = os.getenv("REPORTS_DIRECTORY")
        
    observer = Observer()
    observer.schedule(Main(), DATA_DIRECTORY, recursive = True)
    observer.start()

    logging.info(f"observation has begun")

    try:
        while observer.is_alive():
            observer.join(1)
    finally:
        observer.stop()
        observer.join()